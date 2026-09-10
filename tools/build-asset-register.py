"""
Turn the AFOSI asset register spreadsheet into the JSON the public asset page
reads (public/data/assets.json).

Run this again whenever the register changes:

    python tools/build-asset-register.py "C:/path/to/Asset Register updated.xlsx"

Deliberately NOT published: the purchase amount, the depreciation column and the
internal comments. The QR tag URLs are public and guessable (AFOSI-001,
AFOSI-002 ...), so anything written here should be safe for a stranger to read.
Money stays in the spreadsheet.
"""

import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = Path.home() / "Downloads" / "Asset Register updated.xlsx"
OUT = ROOT / "public" / "data" / "assets.json"

# Public columns only. Amount (D), depreciation (E) and comments (I) are read
# but never written out.
COL = {"tag": 0, "name": 1, "description": 2, "custodian": 5, "condition": 6, "donor": 7}

CATEGORIES = [
    ("Laptop", ("laptop", "thinkpad", "elitebook", "probook", "notebook", "legion",
                "pavillion", "pavilion", "specture", "spectre", "macbook")),
    ("Furniture", ("chair", "desk", "table", "cabinet", "shelf", "cupboard", "sofa")),
    ("Office equipment", ("photocopier", "printer", "scanner", "shredder", "projector",
                          "whiteboard", "camera")),
    ("Appliance", ("microwave", "vacuum", "vaccum", "kettle", "fridge", "refrigerator",
                   "cooker", "dispenser", "fan")),
]


def clean(value):
    """Trim, collapse whitespace and repair the mojibake in the source file."""
    if value is None:
        return ""
    text = str(value)
    # The register was typed with smart quotes for screen sizes ("14.0" Intel...")
    # which came through the export as U+FFFD.
    text = re.sub(r'(\d)\ufffd', r'\1" ', text)
    text = text.replace("\ufffd", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def slug(tag):
    """AFOSI/015/Childfund -> AFOSI-015. The number is unique, so the QR URL
    stays short and every variant of the printed tag resolves to it."""
    match = re.search(r"(\d{2,})", tag)
    return f"AFOSI-{match.group(1)}" if match else re.sub(r"[^A-Za-z0-9]+", "-", tag).strip("-").upper()


def categorise(name, description):
    haystack = f"{name} {description}".lower()
    for label, words in CATEGORIES:
        if any(word in haystack for word in words):
            return label
    return "Equipment"


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not src.exists():
        sys.exit(f"Spreadsheet not found: {src}")

    ws = openpyxl.load_workbook(src, data_only=True).active
    assets, seen = [], set()

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=9, values_only=True):
        tag = clean(row[COL["tag"]])
        name = clean(row[COL["name"]])
        if not tag or not name:
            continue
        asset_id = slug(tag)
        if asset_id in seen:
            sys.exit(f"Duplicate asset id {asset_id} — two rows map to the same QR URL.")
        seen.add(asset_id)

        condition = clean(row[COL["condition"]]).title() or "Unknown"
        assets.append({
            "id": asset_id,
            "tag": tag,
            "name": name,
            "description": clean(row[COL["description"]]),
            "category": categorise(name, clean(row[COL["description"]])),
            "custodian": clean(row[COL["custodian"]]) or "AFOSI office",
            "condition": condition,
            "donor": clean(row[COL["donor"]]) or "AFOSI",
        })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps({"updated": date.today().isoformat(), "count": len(assets), "assets": assets},
                   indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"{len(assets)} assets -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
